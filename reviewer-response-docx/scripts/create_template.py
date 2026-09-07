#!/usr/bin/env python3
"""Create the sanitized Word template distributed with this skill."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor

out=Path(__file__).resolve().parents[1]/"assets/response-template.docx"
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2.54);s.left_margin=s.right_margin=Cm(3.175)
st=d.styles["Normal"];st.font.name="Calibri";st.font.size=Pt(12);st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_after=Pt(6)
def p(text,color=None,bold=False,italic=False):
    x=d.add_paragraph();r=x.add_run(text);r.bold=bold;r.italic=italic
    if color:r.font.color.rgb=RGBColor.from_string(color)
p("Response to Reviewers",bold=True);p("Manuscript reference: [ID]");p("Title: [TITLE]")
p("WORKING DRAFT — replace or remove this status line","C00000",True)
p("Reviewer comments are reproduced in blue, responses in black, and revised manuscript excerpts in red italics. Locations refer to the final clean Word manuscript.")
p("Reviewer 1","255CDB",True);p("Comment 1","255CDB",True);p("[VERBATIM REVIEWER COMMENT]","255CDB")
p("Response",bold=True);p("[DIRECT RESPONSE WITH EVIDENCE AND BOUNDARIES]")
p("Revision","C00000",True);p("(Section [X]; Manuscript, pp. [X–Y], paragraph [N])",bold=True);p("[VERBATIM REVISED MANUSCRIPT EXCERPT]","C00000",italic=True)
d.save(out);print(out)
